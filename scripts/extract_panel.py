#!/usr/bin/env python3
"""Build a quarterly panel from every Supplement P edition in data/supplements/.

extract_sda.py reads one edition into the shape the explorer renders. This
reads all of them into one tidy long-format table -- quarter x geography x
design category x measure -- so the same region and category can be followed
over time, and writes a validation report that says, edition by edition, how
far each quarter can be trusted.

The editions are not laid out alike. Table numbers, sheet names, column
headings and whole tables come and go, so nothing here is keyed on a table
number: each sheet is identified by what its caption and header say it holds
(`classify`), which yields an explicit per-edition table map. A heading the
map has not seen before stops the build rather than being guessed at.

Usage:  python3 scripts/extract_panel.py [data/supplements] [-o data/panel]

Writes, into the output directory:
  panel.csv        the panel, one row per published or derived figure
  figure_p1.csv    Figure P.1's national history as each edition states it
  validation.json  per-edition table map, checks and detected breaks
  VALIDATION.md    the same, written out for reading
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from extract_sda import (  # noqa: E402
    DESIGN_CATEGORIES,
    DWELLING_LABEL,
    RESIDENT_COUNTS,
    extract_national_trend,
    parse_value,
    to_transitional,
)

MONTHS = {m: i for i, m in enumerate(
    "January February March April May June July August September October "
    "November December".split(), 1)}
QUARTER_END = {3: 31, 6: 30, 9: 30, 12: 31}
STATES = ["ACT", "NSW", "NT", "QLD", "SA", "TAS", "VIC", "WA"]

# A run of this many empty rows ends a sheet. One 2024-25 Q3 worksheet is
# styled across all 16,384 columns and expands to 85 MB of XML; bounding the
# width by the header and stopping at the last populated row keeps every
# edition to a few seconds.
BLANK_RUN = 50


# --------------------------------------------------------------------------
# The table map
# --------------------------------------------------------------------------

def _norm(text) -> str:
    return " ".join(str(text or "").lower().split())


def classify(caption: str, header: list[str]):
    """Name the table a sheet holds, from its caption and header alone.

    Returns (role, level, scheme), or None for sheets that are not tables.
    `scheme` distinguishes versions of the same table whose columns mean
    different things, so a definitional change is visible in the map rather
    than hidden in a column rename.
    """
    c = _norm(caption)
    h = [_norm(x) for x in header]
    if not c.startswith("table p."):
        return None
    first = h[0] if h else ""
    if first == "service district":
        return "participants_by_service_district", "Service District", None
    if first == "state/territory":
        role = "providers_by_state" if "providers" in c else "committed_supports_by_state"
        return role, "State", None

    level = "SA3" if first.startswith("sa3") else "SA4"
    joined = " | ".join(h)

    # Every stock table until 2025-26 Q1 is captioned "(excluding in-kind
    # arrangements)"; only the tables *of* in-kind dwellings are in-kind.
    if "in-kind" in c.replace("excluding in-kind", ""):
        if "maximum number of residents" in c:
            return "inkind_max_residents", level, None
        if "build type" in c:
            return "inkind_existing_legacy_detail", level, None
        return "inkind_categories", level, None
    if "by status" in c:
        if "legacy crm" in joined:
            scheme = "legacy_crm"
        elif "sda funding, sda in use" in joined:
            scheme = "funding_split"
        else:
            scheme = "in_use_eligible"
        return "demand_status", level, scheme
    # 2023-24 Q2 and Q3 caption their SA3 not-in-use table as the old
    # "seeking SDA" table; the header is what it actually holds.
    if "sda not in use" in joined and "percentage" in joined:
        return "not_using", level, None
    if "seeking sda" in c or "sda decision design category" in c:
        if "not defined" in joined:
            scheme = "seeking_legacy"
        elif "basic" in h:
            scheme = "eligible_decision_with_basic"
        else:
            scheme = "eligible_decision"
        return "demand_categories", level, scheme
    if "unfinished" in c:
        return ("pipeline_detail" if "build type" in c else "pipeline_categories"), level, None
    if "maximum residents by design category" in c:
        return "newbuild_places", level, None
    if "maximum number of residents" in c:
        return "max_residents", level, None
    if "existing stock and legacy stock" in c:
        return "existing_legacy_detail", level, None
    if "new build and new build (refurbished) dwellings" in c:
        return "newbuild_detail", level, None
    # SA4 captions say "Building Type" and SA3 captions "Build Type".
    if "building type" in c or ("build type" in c and "design category" not in c):
        return "build_types", level, None
    if "design category" in c and "enrolled" in c:
        return "stock_categories", level, None
    raise ValueError(f"unrecognised table: {caption!r}")


def edition_label(path: Path) -> str:
    """'Supplement_P_SDA_2025-26_Q4.xlsx' -> '2025-26 Q4'."""
    match = re.search(r"(\d{4}-\d{2})_Q(\d)", path.name)
    if not match:
        raise ValueError(f"cannot read an edition from {path.name}")
    return f"{match.group(1)} Q{match.group(2)}"


def as_at_date(caption: str):
    """'... as at 30 June 2026 (excluding ...)' -> '2026-06-30'."""
    match = re.search(r"as at (\d{1,2}) (\w+) (\d{4})", caption)
    if not match:
        return None
    day, month, year = match.groups()
    return date(int(year), MONTHS[month], int(day)).isoformat()


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------

def table_width(header_row) -> int:
    """Columns up to the first gap or Excel's 'ColumnN' placeholder.

    The existing/legacy cross-tab is an Excel table stretched to 16,384
    columns, 16,322 of them headed Column1, Column2, ... and empty.
    """
    width = 0
    for i, cell in enumerate(header_row):
        if cell is None or (i and re.fullmatch(r"Column\d+", str(cell).strip())):
            break
        width = i + 1
    return width


def read_header(ws):
    rows = list(ws.iter_rows(values_only=True, max_row=2))
    caption = str(rows[0][0]).strip() if rows and rows[0] and rows[0][0] else ""
    header_row = rows[1] if len(rows) > 1 else ()
    width = table_width(header_row)
    header = [str(h).strip() for h in header_row[:width]]
    return caption, header


def read_body(ws, header):
    """Data rows as {label: {"values", "flags"}}, plus notes and duplicates.

    Row 1 is the caption, row 2 the header, then data rows and footnotes.
    Footnotes are long prose in the label column, so length separates them
    from region names, as in extract_sda.read_table. A label that repeats is
    dropped if the repeat is identical and fails the build if not.
    """
    width = len(header)
    columns = header[1:]
    rows, notes, duplicates = {}, [], []
    blank = 0
    for raw in ws.iter_rows(values_only=True, min_row=3, max_col=width):
        if all(v is None for v in raw):
            blank += 1
            if blank >= BLANK_RUN:
                break
            continue
        blank = 0
        label = raw[0]
        if label is None:
            continue
        label = str(label).strip()
        if label.startswith(("Back to", "Go to")):
            continue
        if len(label) > 60:
            notes.append(label)
            continue
        record, flags = {}, {}
        for column, cell in zip(columns, raw[1:]):
            value, flag = parse_value(cell)
            record[column] = value
            if flag:
                flags[column] = flag
        entry = {"values": record, "flags": flags}
        if label in rows:
            if rows[label] != entry:
                raise ValueError(f"{ws.title}: conflicting duplicate rows for {label!r}")
            duplicates.append(label)
            continue
        rows[label] = entry
    return rows, notes, duplicates


def read_edition(path: Path, workdir: Path, level: str = "SA4"):
    """Open one edition and read every table the panel needs at one level.

    `level` is "SA4" for the main panel or "SA3" for extract_panel_sa3.py. The
    SA3 read skips Figure P.1's chart, which only the SA4 report checks.
    """
    import openpyxl

    converted = to_transitional(path, workdir / (path.stem + ".xlsx"))
    workbook = openpyxl.load_workbook(converted, read_only=True, data_only=True)

    table_map, tables = [], {}
    as_at = None
    for ws in workbook.worksheets:
        caption, header = read_header(ws)
        found = classify(caption, header)
        if not found:
            continue
        role, table_level, scheme = found
        entry = {"sheet": ws.title, "role": role, "level": table_level, "scheme": scheme,
                 "caption": caption, "columns": header[1:]}
        table_map.append(entry)
        if role == "stock_categories" and table_level == level:
            as_at = as_at_date(caption)
        if table_level != level:
            continue
        if role in tables:
            raise ValueError(f"{path.name}: two {level} tables classified as {role}")
        rows, notes, duplicates = read_body(ws, header)
        tables[role] = {**entry, "rows": rows, "notes": notes, "duplicates": duplicates}

    figure_prose = [str(cell) for row in workbook["Figure P.1"].iter_rows(values_only=True)
                    for cell in row if cell]
    intro = [str(cell) for row in workbook["Intro"].iter_rows(values_only=True, max_col=3)
             for cell in row if cell]
    workbook.close()
    converted.unlink()
    return {
        "edition": edition_label(path),
        "source": path.name,
        "as_at": as_at,
        "table_map": table_map,
        "tables": tables,
        "figure_prose": figure_prose,
        "intro": intro,
        "chart_trend": extract_national_trend(path, ()) if level == "SA4" else None,
    }


# --------------------------------------------------------------------------
# Figure P.1
# --------------------------------------------------------------------------

FIGURE_SENTENCES = [
    # (pattern after "there were N ", series)
    (r"active participants with SDA in use", "participants_sda_in_use"),
    (r"active participants with SDA supports", "participants_with_sda_supports"),
    (r"active participants with SIL supports", "participants_with_sil"),
    (r"in annualised SDA supports", "sda_annualised_m"),
    (r"in annualised SIL supports", "sil_annualised_m"),
    (r"enrolled dwellings", "enrolled_dwellings"),
]


def quarter_end(month: int, year: int) -> str:
    return date(year, month, QUARTER_END[month]).isoformat()


def parse_figure_prose(texts):
    """Figure P.1's history from its accessibility description.

    The prose is present in every edition, where the chart XML exists only
    from 2024-25 Q2 and carries filtered duplicate series. Each sentence reads
    "In <Month> <Year> there were <N> <what>", and from 2024-25 Q1 the
    participant sentence also carries "and <M> active participants eligible
    but not yet using SDA".
    """
    out = {}
    sentence = re.compile(r"In (\w+) (\d{4}) there (?:were|was) \$?([\d,]+)m? ([^.\n]+)")
    for text in texts:
        for month, year, number, rest in sentence.findall(text):
            if month not in MONTHS:
                continue
            quarter = quarter_end(MONTHS[month], int(year))
            for pattern, series in FIGURE_SENTENCES:
                if rest.startswith(pattern):
                    out[(series, quarter)] = float(number.replace(",", ""))
                    break
            also = re.search(r"and ([\d,]+) active participants eligible but not yet using", rest)
            if also:
                out[("participants_eligible_not_using", quarter)] = float(
                    also.group(1).replace(",", ""))
    return out


def chart_series(chart_trend):
    """The chart XML's series, keyed like parse_figure_prose's output."""
    names = {"sda_in_use": "participants_sda_in_use",
             "sda_eligible_not_using": "participants_eligible_not_using",
             "sil_participants": "participants_with_sil",
             "sda_annualised": "sda_annualised_m",
             "sil_annualised": "sil_annualised_m",
             "enrolled_dwellings": "enrolled_dwellings"}
    months = {m[:3]: i for m, i in MONTHS.items()}
    out = {}
    quarters = chart_trend.get("quarters") or []
    for key, values in (chart_trend.get("series") or {}).items():
        if key not in names or len(values) != len(quarters):
            continue
        for label, value in zip(quarters, values):
            mon, yy = label.split("-")
            if value is None:
                continue
            # The chart carries dollars where the prose (and the series
            # label) says $m.
            if names[key].endswith("_m"):
                value = round(value / 1e6)
            out[(names[key], quarter_end(months[mon], 2000 + int(yy)))] = value
    return out


# --------------------------------------------------------------------------
# From tables to panel records
# --------------------------------------------------------------------------

def geography(label: str, level: str = "SA4"):
    """(id, level, state) for a row label, matching sda.json's ids.

    '<STATE> - Other' rows appear only in the SA4 participant tables and hold
    participants the NDIA could not place in an SA4; 'Missing' holds those
    it could not place in a state. Neither is a region. SA3 tables carry an
    '- Other' row in every table, and it is zero in the stock tables.
    """
    if label == "Total":
        return "national", "National", None
    if label in STATES:
        return f"state:{label}", "State", label
    if label == "Missing":
        return "unknown", "Unknown", None
    if " - " in label:
        state, name = label.split(" - ", 1)
        if state not in STATES:
            raise ValueError(f"unrecognised state in {label!r}")
        if name == "Other":
            return f"other:{state}", "Other", state
        return f"{level.lower()}:{label}", level, state
    raise ValueError(f"unrecognised geography {label!r}")


STATUS_COLUMNS = {
    "participants in an sda dwelling and no evidence of seeking an alternative (legacy crm data)":
        "legacy_in_sda_not_seeking",
    "participants in an sda dwelling, seeking alternative (legacy crm data)":
        "legacy_in_sda_seeking_alternative",
    "additional participants eligible for sda (legacy crm data)": "legacy_eligible_unfunded",
    "participants with sda funding": "participants_sda_funded",
    "participants with sda funding, sda in use": "participants_sda_in_use",
    "participants with sda funding, sda not in use": "participants_funded_not_in_use",
    "additional participants eligible for sda": "participants_eligible_unfunded",
    "participants with sda in use": "participants_sda_in_use",
    "participants sda eligible, not yet using sda": "participants_eligible_not_using",
    "total participants with sda funding or an sda need": "participants_total_need",
    "total participants with sda in use or sda eligible (not yet using sda)":
        "participants_total_need",
}

DEMAND_COLUMNS = {
    "not defined": "Not Defined",
    "missing": "Missing",
    "total participants seeking sda (legacy crm data)": "Total",
    "total participants with an sda need": "Total",
    "percentage of participants seeking sda dwellings": None,
}

BUILD_TYPE_COLUMNS = {
    "existing": "dwellings_existing",
    "legacy": "dwellings_legacy",
    "new build": "dwellings_new_build",
    "new build (refurbished)": "dwellings_new_build_refurbished",
    # The SA3 table heads the same column "New Build (Refurbishment)".
    "new build (refurbishment)": "dwellings_new_build_refurbished",
    "total": None,
}

MAX_RESIDENT_MEASURES = {label: f"dwellings_max_{n}{'plus' if n == 6 else ''}_resident"
                         f"{'' if n == 1 else 's'}" for label, n in RESIDENT_COUNTS.items()}

CATEGORY_ORDER = DESIGN_CATEGORIES + ["Not Defined", "Missing", "Total"]


def category_column(column: str):
    """A design-category column heading, or 'Total', or None."""
    if column in DESIGN_CATEGORIES:
        return column
    if _norm(column) == "total":
        return "Total"
    return None


def crosstab(values, flags):
    """Collapse a dwelling-form cross-tab row into dwellings and places by category.

    Each heading names a dwelling form and its resident count, so dwellings x
    residents is places. A category any of whose cells is suppressed comes
    back as None rather than as the sum of the cells that were published.
    """
    out = {}
    for label, value in values.items():
        if _norm(label) == "total":
            continue
        match = DWELLING_LABEL.match(label)
        if not match:
            raise ValueError(f"unparsed cross-tab heading {label!r}")
        residents = int(match.group("residents") or match.group("first"))
        category = match.group("category").strip()
        cell = out.setdefault(category, {"dwellings": 0.0, "places": 0.0, "flag": None})
        if value is None:
            cell["flag"] = flags.get(label, "missing")
            continue
        cell["dwellings"] += value
        cell["places"] += value * residents
    for cell in out.values():
        if cell["flag"]:
            cell["dwellings"] = cell["places"] = None
    return out


class Panel:
    """Records for one edition, keyed (geography, category, measure)."""

    def __init__(self, level="SA4"):
        self.records = {}
        self.geos = {}
        self.level = level

    def put(self, label_or_geo, category, measure, value, flag, source):
        if isinstance(label_or_geo, tuple):
            geo = label_or_geo
        else:
            geo = geography(label_or_geo, self.level)
        self.geos[geo[0]] = geo
        self.records[(geo[0], category, measure)] = (value, flag or "", source)

    def get(self, geo_id, category, measure):
        found = self.records.get((geo_id, category, measure))
        return found[0] if found else None

    def has(self, geo_id, category, measure):
        return (geo_id, category, measure) in self.records


def build_panel(edition, level="SA4"):
    """Turn one edition's tables into panel records.

    Published figures keep the table they came from in `source`. Derived ones
    say how they were derived, and follow extract_sda's conventions exactly:
    places are P.7's published new-build places plus existing/legacy
    dwellings x residents from P.12, and places not in SDA use is enrolled
    places less participants with SDA in use.

    At SA3 neither P.7 nor the cross-tabs are published, so enrolled places
    (by category or in total) and places not in SDA use cannot be derived;
    places from dwellings by maximum residents is the only places measure.
    """
    tables = edition["tables"]
    panel = Panel(level)

    def sheet(role):
        return tables[role]["sheet"].replace("Table ", "")

    def simple(role, measure_for):
        if role not in tables:
            return
        src = sheet(role)
        for label, row in tables[role]["rows"].items():
            for column, value in row["values"].items():
                target = measure_for(column)
                if target is None:
                    continue
                category, measure = target
                panel.put(label, category, measure, value, row["flags"].get(column), src)

    def by_category(measure):
        def pick(column):
            category = category_column(column)
            return (category, measure) if category else None
        return pick

    simple("stock_categories", by_category("enrolled_dwellings"))
    simple("newbuild_places", by_category("newbuild_places"))
    simple("pipeline_categories", by_category("pipeline_dwellings"))
    simple("inkind_categories", by_category("inkind_dwellings"))
    simple("build_types", lambda c: ("Total", BUILD_TYPE_COLUMNS[_norm(c)])
           if BUILD_TYPE_COLUMNS[_norm(c)] else None)
    simple("max_residents", lambda c: ("Total", MAX_RESIDENT_MEASURES[c])
           if c in MAX_RESIDENT_MEASURES else None)

    def status_measure(column):
        key = _norm(column)
        if key not in STATUS_COLUMNS:
            raise ValueError(f"{edition['edition']}: unmapped status column {column!r}")
        measure = STATUS_COLUMNS[key]
        # Same heading, different content: under the legacy CRM scheme the
        # total adds a legacy count of "additional eligible" participants,
        # and it sits 2-3% above the later series for the same quarter.
        if measure == "participants_total_need" and \
                tables["demand_status"]["scheme"] == "legacy_crm":
            measure = "legacy_total_funded_or_eligible"
        return "Total", measure
    simple("demand_status", status_measure)

    if "demand_categories" in tables:
        scheme = tables["demand_categories"]["scheme"]
        measure = ("legacy_participants_seeking" if scheme == "seeking_legacy"
                   else "participants_with_need")

        def demand_measure(column):
            category = category_column(column)
            if category is None:
                key = _norm(column)
                if key not in DEMAND_COLUMNS:
                    raise ValueError(f"{edition['edition']}: unmapped demand column {column!r}")
                category = DEMAND_COLUMNS[key]
            return (category, measure) if category else None
        simple("demand_categories", demand_measure)

    # Cross-tabs: dwellings and derived places by category.
    detail = {}
    for role, dwell_measure, place_measure in (
            ("newbuild_detail", "newbuild_dwellings", None),
            ("existing_legacy_detail", "existing_legacy_dwellings", "existing_legacy_places"),
            ("pipeline_detail", None, "pipeline_places")):
        if role not in tables:
            continue
        src = sheet(role)
        for label, row in tables[role]["rows"].items():
            cells = crosstab(row["values"], row["flags"])
            detail[(role, label)] = cells
            for category, cell in cells.items():
                if dwell_measure:
                    panel.put(label, category, dwell_measure, cell["dwellings"],
                              cell["flag"], f"{src} sum")
                if place_measure:
                    panel.put(label, category, place_measure, cell["places"],
                              cell["flag"], f"{src} x residents")
            if place_measure:
                parts = [c["places"] for c in cells.values()]
                total = None if any(p is None for p in parts) else sum(parts)
                panel.put(label, "Total", place_measure, total,
                          None if total is not None else "incomplete",
                          f"{src} x residents")

    # Enrolled places: published new-build places plus derived existing/legacy.
    if "newbuild_detail" in tables and "existing_legacy_detail" in tables:
        p7 = tables.get("newbuild_places")
        corrections = []
        for label in tables["existing_legacy_detail"]["rows"]:
            geo = geography(label, level)
            newbuild = detail.get(("newbuild_detail", label), {})
            existing = detail[("existing_legacy_detail", label)]
            total, complete, corrected = 0.0, True, None
            for category in DESIGN_CATEGORIES:
                published = (p7["rows"].get(label, {}).get("values", {}).get(category)
                             if p7 and category in p7["columns"] else None)
                built = (newbuild.get(category) or {}).get("dwellings") or 0
                flag = None
                if published == 0 and built > 0:
                    # P.7 occasionally publishes zero places against new-build
                    # dwellings P.11 lists in the same region and category
                    # (Wheat Belt, 2025-26 Q1-Q3). Zero places cannot hold a
                    # dwelling, so the derivation stands in, flagged.
                    published, flag = None, "p7_zero_with_dwellings"
                    corrected = flag
                if published is not None:
                    new_pl, new_src = published, sheet("newbuild_places")
                else:
                    # P.7 has no Basic column (and would fall back here if a
                    # cell were ever suppressed), so the cross-tab stands in.
                    # Basic new builds are zero in every edition so far.
                    new_pl = (newbuild.get(category) or {"places": 0.0})["places"]
                    new_src = f"{sheet('newbuild_detail')} x residents"
                    if flag and geo[1] == "SA4":
                        corrections.append((geo, category, new_pl))
                old_pl = (existing.get(category) or {"places": 0.0})["places"]
                if new_pl is None or old_pl is None:
                    places, complete = None, False
                else:
                    places = new_pl + old_pl
                    total += places
                panel.put(geo, category, "enrolled_places", places,
                          flag if places is not None else "incomplete",
                          f"{new_src} + {sheet('existing_legacy_detail')} x residents")
            panel.put(geo, "Total", "enrolled_places", total if complete else None,
                      corrected if complete else "incomplete", "sum of categories")

        # P.7's state and national subtotals omit the same places, so a region
        # corrected above is carried up to them, or they would not reconcile.
        for geo, category, delta in corrections:
            for parent in (f"state:{geo[2]}", "national"):
                for cat in (category, "Total"):
                    old, _, src = panel.records[(parent, cat, "enrolled_places")]
                    if old is not None:
                        panel.records[(parent, cat, "enrolled_places")] = (
                            old + delta, "p7_zero_with_dwellings", src)

    # Places implied by dwellings by maximum residents (P.6): an independent
    # check on the derivation, covering every build type at once.
    if "max_residents" in tables:
        src = sheet("max_residents")
        for label, row in tables["max_residents"]["rows"].items():
            vals = row["values"]
            if any(vals.get(k) is None for k in RESIDENT_COUNTS):
                places = None
            else:
                places = sum(vals[k] * n for k, n in RESIDENT_COUNTS.items())
            panel.put(label, "Total", "places_from_max_residents", places,
                      None if places is not None else "incomplete",
                      f"{src} x residents")

    # Harmonise the three status schemes where the concepts line up.
    if "demand_status" in tables:
        scheme = tables["demand_status"]["scheme"]
        src = sheet("demand_status")
        for geo_id, geo in list(panel.geos.items()):
            if scheme == "funding_split":
                a = panel.get(geo_id, "Total", "participants_funded_not_in_use")
                b = panel.get(geo_id, "Total", "participants_eligible_unfunded")
                if panel.has(geo_id, "Total", "participants_funded_not_in_use"):
                    panel.put(geo, "Total", "participants_eligible_not_using",
                              a + b if a is not None and b is not None else None,
                              "harmonised",
                              f"{src} not in use + additional eligible")

    # 2024-25 Q3's participant figures are as at 2 April, not 31 March.
    if any("as of 2 april" in _norm(t) for t in edition["intro"]):
        for key, (value, flag, source) in panel.records.items():
            if key[2].startswith(("participants", "legacy")) and not flag:
                panel.records[key] = (value, "participants_as_at_2_april", source)

    # Places not in SDA use.
    for geo_id, geo in list(panel.geos.items()):
        places = panel.get(geo_id, "Total", "enrolled_places")
        in_use = panel.get(geo_id, "Total", "participants_sda_in_use")
        if places is not None and in_use is not None:
            panel.put(geo, "Total", "places_not_in_use", places - in_use, None,
                      "places - in use")

    return panel


# --------------------------------------------------------------------------
# Output
# --------------------------------------------------------------------------

LEVEL_ORDER = {"National": 0, "State": 1, "SA4": 2, "Other": 3, "Unknown": 4}
PANEL_COLUMNS = ["as_at", "geography", "level", "state", "category",
                 "measure", "value", "flag", "source"]


def fmt(value):
    if value is None:
        return ""
    if float(value).is_integer():
        return str(int(value))
    return f"{value:.4f}".rstrip("0")


def panel_rows(edition, panel):
    rows = []
    for (geo_id, category, measure), (value, flag, source) in panel.records.items():
        _, level, state = panel.geos[geo_id]
        rows.append([edition["as_at"], geo_id, level, state or "",
                     category, measure, fmt(value), flag, source])
    rows.sort(key=lambda r: (r[0], LEVEL_ORDER[r[2]], r[1],
                             CATEGORY_ORDER.index(r[4]), r[5]))
    return rows


def write_csv(path: Path, header, rows):
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    path.write_text(buffer.getvalue())


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("supplements", type=Path, nargs="?", default=Path("data/supplements"))
    parser.add_argument("-o", "--out", type=Path, default=Path("data/panel"))
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    import validate_panel  # noqa: E402  (sibling module, kept apart for size)

    paths = sorted(p for p in args.supplements.iterdir()
                   if re.search(r"Supplement_P_SDA_\d{4}-\d{2}_Q\d\.xls[xb]$", p.name))
    editions, panels = [], []
    with tempfile.TemporaryDirectory() as tmp:
        for path in paths:
            edition = read_edition(path, Path(tmp))
            edition["figure"] = parse_figure_prose(edition["figure_prose"])
            edition["chart"] = chart_series(edition.pop("chart_trend"))
            editions.append(edition)
            panels.append(build_panel(edition))
            print(f"  read {edition['edition']} (as at {edition['as_at']})", flush=True)

    rows = [r for e, p in zip(editions, panels) for r in panel_rows(e, p)]
    write_csv(args.out / "panel.csv", PANEL_COLUMNS, rows)

    figure_rows = sorted(
        [e["edition"], e["as_at"], quarter, series, fmt(value)]
        for e in editions for (series, quarter), value in e["figure"].items())
    write_csv(args.out / "figure_p1.csv",
              ["edition", "edition_as_at", "quarter", "series", "value"], figure_rows)

    report = validate_panel.validate(editions, panels)
    (args.out / "validation.json").write_text(json.dumps(report, indent=1) + "\n")
    (args.out / "VALIDATION.md").write_text(validate_panel.to_markdown(report))

    print(f"wrote {args.out}/panel.csv  ({len(rows):,} rows, "
          f"{(args.out / 'panel.csv').stat().st_size / 1e6:.1f} MB)")
    validate_panel.print_summary(report)


if __name__ == "__main__":
    main()
